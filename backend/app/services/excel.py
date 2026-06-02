import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.services.hours import build_daily_summaries

# Paleta
HEADER_FILL = PatternFill(start_color="375623", end_color="375623", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
EVEN_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
OT_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # naranja claro
INCOMPLETE_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # amarillo
DATA_FONT = Font(name="Calibri", size=11, color="000000")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def _style_header(ws, n_cols):
    for col_idx in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = THIN_BORDER


def _autofit(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if isinstance(cell.value, datetime):
                cell_len = 18
            elif cell.value is not None:
                cell_len = len(str(cell.value))
            else:
                cell_len = 0
            max_len = max(max_len, cell_len)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)


def _fmt_time(iso_str):
    if not iso_str:
        return "—"
    return datetime.fromisoformat(iso_str).strftime("%d/%m/%Y %H:%M")


def generate_attendance_excel(logs, sheet_name="Asistencia"):
    """Genera un Excel con dos hojas:
    1) 'Resumen de Horas': por empleado/día con entrada, salida, horas trabajadas y extra.
    2) 'Marcaciones': bitácora cronológica detallada de cada marca.
    """
    wb = Workbook()

    # ---------- Hoja 1: Resumen de Horas ----------
    ws = wb.active
    ws.title = "Resumen de Horas"
    headers = [
        "Person ID", "Nombre", "Departamento", "Turno", "Fecha",
        "Entrada", "Salida", "Horas Trabajadas", "Horas Extra", "Estado",
    ]
    ws.append(headers)
    _style_header(ws, len(headers))

    summaries = build_daily_summaries(logs)
    total_worked = 0.0
    total_overtime = 0.0

    for r_idx, s in enumerate(summaries, start=2):
        estado = "Incompleto" if s["incomplete"] else "Completo"
        ws.append([
            f"'{s['person_id']}",
            s["name"].upper(),
            s["department"].upper(),
            s["shift_name"],
            datetime.fromisoformat(s["date"]).strftime("%d/%m/%Y"),
            _fmt_time(s["first_in"]),
            _fmt_time(s["last_out"]),
            s["worked_hours"],
            s["overtime_hours"],
            estado,
        ])
        total_worked += s["worked_hours"]
        total_overtime += s["overtime_hours"]

        if s["incomplete"]:
            row_fill = INCOMPLETE_FILL
        elif s["overtime_hours"] > 0:
            row_fill = OT_FILL
        else:
            row_fill = WHITE_FILL if r_idx % 2 == 0 else EVEN_FILL

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.fill = row_fill
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # Fila de totales
    if summaries:
        total_row = len(summaries) + 2
        ws.cell(row=total_row, column=7, value="TOTALES:")
        ws.cell(row=total_row, column=8, value=round(total_worked, 2))
        ws.cell(row=total_row, column=9, value=round(total_overtime, 2))
        for col_idx in range(7, 10):
            cell = ws.cell(row=total_row, column=col_idx)
            cell.font = Font(name="Calibri", size=11, bold=True, color="000000")
            cell.border = THIN_BORDER
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(summaries) + 1}"

    _autofit(ws)

    # ---------- Hoja 2: Marcaciones (bitácora) ----------
    ws2 = wb.create_sheet(title="Marcaciones")
    headers2 = ["Person ID", "Nombre", "Departamento", "Tipo", "Fecha y Hora"]
    ws2.append(headers2)
    _style_header(ws2, len(headers2))

    sorted_logs = sorted(logs, key=lambda l: l.timestamp)
    for r_idx, log in enumerate(sorted_logs, start=2):
        name_val = log.employee.name.upper() if log.employee else "DESCONOCIDO"
        dept_val = log.employee.department.upper() if log.employee else "SIN DEPARTAMENTO"
        ws2.append([
            f"'{log.person_id}",
            name_val,
            dept_val,
            (log.event_type or "entrada").capitalize(),
            log.timestamp,
        ])
        row_fill = WHITE_FILL if r_idx % 2 == 0 else EVEN_FILL
        for col_idx in range(1, len(headers2) + 1):
            cell = ws2.cell(row=r_idx, column=col_idx)
            cell.fill = row_fill
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center")
            if col_idx == 5:
                cell.number_format = "d/m/yyyy h:mm"

    if sorted_logs:
        ws2.auto_filter.ref = f"A1:{get_column_letter(len(headers2))}{len(sorted_logs) + 1}"
    _autofit(ws2)

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    return file_stream.getvalue()
